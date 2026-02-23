import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook


def excel_safe(value):
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    try:
        return json.dumps(value, ensure_ascii=False)
    except Exception:
        return str(value)


def write_table(ws, rows, headers=None):
    if not rows:
        return

    if headers is None:
        headers = sorted({k for r in rows for k in r})

    ws.append(headers)
    for r in rows:
        ws.append([excel_safe(r.get(h)) for h in headers])


def write_kv(ws, d):
    ws.append(["key", "value"])
    for k, v in d.items():
        ws.append([k, excel_safe(v)])


def write_dataframe(ws, df: pd.DataFrame):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([excel_safe(v) for v in row])


def load_jsonl_records(path: Path):
    records = []
    with path.open("r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                records.append(json.loads(line))
            except Exception:
                continue
    return records


def flatten_stage_record(stage_name, record):
    flat = {"stage": stage_name}
    for key, value in record.items():
        if isinstance(value, dict):
            for sub_key, sub_val in value.items():
                flat[f"{key}.{sub_key}"] = sub_val
        else:
            flat[key] = value
    return flat


def add_dataset_sheet(wb, sheet_name, csv_path: Path, max_rows=2000):
    if not csv_path or not csv_path.exists():
        return

    ws = wb.create_sheet(sheet_name)
    try:
        df = pd.read_csv(csv_path)
    except Exception:
        ws.append(["error", f"failed to read {csv_path}"])
        return

    if len(df) > max_rows:
        df = df.head(max_rows)
        ws.append(["note", f"showing first {max_rows} rows only"])

    write_dataframe(ws, df)


def export_full_audit(data_dir, output_file):
    data_dir = Path(data_dir)

    wb = Workbook()
    wb.remove(wb.active)

    dataset_meta = {}
    column_profiles = []
    preprocess_steps = []
    stage_rollup = []
    artifacts = []

    stage_names = {
        "model_selection",
        "model_initialization",
        "ml_experiments",
        "ml_experimentation",
        "evaluation",
    }

    for jf in data_dir.glob("*.jsonl"):
        lines = load_jsonl_records(jf)
        if not lines:
            continue

        rec = lines[-1]
        name = jf.stem

        if name == "column_inspection":
            dataset_meta.update(
                {
                    "dataset": rec.get("dataset_file_name"),
                    "path": rec.get("dataset_file_path"),
                }
            )
            column_profiles.extend(rec.get("column_profiles", []))

        elif "preprocess" in name:
            preprocess_steps.append(rec)

        elif name in stage_names:
            stage_rollup.append(flatten_stage_record(name, rec))

        else:
            artifacts.append({"artifact": name, **rec})

    if dataset_meta:
        ws = wb.create_sheet("dataset_summary")
        write_kv(ws, dataset_meta)

    if column_profiles:
        ws = wb.create_sheet("column_profiles")
        headers = list(column_profiles[0].keys())
        write_table(ws, column_profiles, headers=headers)

    if preprocess_steps:
        ws = wb.create_sheet("preprocessing")
        write_table(ws, preprocess_steps)

    if stage_rollup:
        ws = wb.create_sheet("pipeline_stages")
        write_table(ws, stage_rollup)

    if artifacts:
        ws = wb.create_sheet("artifacts")
        write_table(ws, artifacts)

    user_input_jsonl = data_dir / "user_input.jsonl"
    if user_input_jsonl.exists():
        user_records = load_jsonl_records(user_input_jsonl)
        if user_records:
            latest_user = user_records[-1]
            upload_path = latest_user.get("data_path") or latest_user.get("dataset_file_path")
            if upload_path:
                add_dataset_sheet(wb, "uploaded_data", Path(upload_path))

    preprocessed_1_files = sorted(
        [p for p in data_dir.glob("*preprocessed_1*.csv") if "final" not in p.name],
        key=lambda p: p.stat().st_mtime,
    )
    if preprocessed_1_files:
        add_dataset_sheet(wb, "preprocessed_1", preprocessed_1_files[-1])

    preprocessed_final_files = sorted(
        [p for p in data_dir.glob("*preprocessed*final*.csv")],
        key=lambda p: p.stat().st_mtime,
    )
    if preprocessed_final_files:
        add_dataset_sheet(wb, "preprocessed_final", preprocessed_final_files[-1])

    wb.save(output_file)
