# OK
import json
from pathlib import Path
from openpyxl import Workbook
from parser.excel_writer import sheet_table


INPUT = Path("data/ml_experiments.jsonl")
OUTPUT = Path("parser/data/ml_experiments.xlsx")


def load(path):
    return [json.loads(l) for l in path.read_text().splitlines() if l.strip()]

def flatten_dic_2(d, parent_key="", sep="_"):
    items = {}

    if isinstance(d, dict):
        for k, v in d.items():
            new_key = f"{parent_key}{sep}{k}" if parent_key else k
            items.update(flatten_dic_2(v, new_key, sep))

    elif isinstance(d, list):
        # store lists as comma string (safe for Excel)
        items[parent_key] = ", ".join(map(str, d))

    else:
        items[parent_key] = d

    return items

def parse_record(wb, rec, i):

    if "experiment_id" not in rec:
        return

    flat = flatten_dic_2(rec)

    sheet = "experiments"

    if sheet not in wb.sheetnames:
        ws = wb.create_sheet(sheet)
        ws.append(list(flat.keys()))
        ws.append(list(flat.values()))
    else:
        ws = wb[sheet]

        # ensure new keys don’t get lost
        existing_headers = [c.value for c in ws[1]]
        new_keys = list(flat.keys())

        if set(new_keys) != set(existing_headers):
            # rebuild header union (important for evolving schema)
            all_keys = sorted(set(existing_headers) | set(new_keys))

            rows = []
            for r in ws.iter_rows(min_row=2, values_only=True):
                row_dict = dict(zip(existing_headers, r))
                rows.append([row_dict.get(k) for k in all_keys])

            ws.delete_rows(1, ws.max_row)
            ws.append(all_keys)
            for r in rows:
                ws.append(r)

        headers = [c.value for c in ws[1]]
        ws.append([flat.get(h) for h in headers])


def main():
    records = load(INPUT)

    wb = Workbook()
    wb.remove(wb.active)

    sheet_count = 0

    for i, rec in enumerate(records, 1):
        before = len(wb.sheetnames)
        parse_record(wb, rec, i)
        after = len(wb.sheetnames)
        if after > before:
            sheet_count += 1

    # ---------- ensure at least one sheet ----------
    if sheet_count == 0:
        ws = wb.create_sheet("no_data")
        ws.append(["message"])
        ws.append(["No experiments found in JSONL"])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)


if __name__ == "__main__":
    main()