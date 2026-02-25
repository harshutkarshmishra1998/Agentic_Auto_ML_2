# OK
import json
from pathlib import Path
from openpyxl import Workbook
from parser.excel_writer import _auto_adjust_column_width


# --------------------------------------------------
# helpers
# --------------------------------------------------
def _fmt(v):
    if v is None:
        return ""
    if isinstance(v, (list, dict)):
        return json.dumps(v)
    return v


def _write_table(ws, rows, headers):
    ws.append(headers)
    for r in rows:
        ws.append([_fmt(r.get(h)) for h in headers])
    _auto_adjust_column_width(ws)


def _write_simple_list(ws, values, header="value"):
    ws.append([header])
    for v in values:
        ws.append([_fmt(v)])
    _auto_adjust_column_width(ws)


def _write_dependency_graph(ws, graph: dict):
    ws.append(["feature", "depends_on"])
    for k, deps in graph.items():
        if not deps:
            ws.append([k, ""])
        else:
            for d in deps:
                ws.append([k, d])
    _auto_adjust_column_width(ws)


# --------------------------------------------------
# main exporter
# --------------------------------------------------
def column_inspection(last_n: int | None = None):
    PROJECT_ROOT = Path(__file__).resolve().parents[1]
    JSONL_PATH = PROJECT_ROOT / "data"
    XLSX_PATH= PROJECT_ROOT / "parser" / "data" / "xlsx"

    XLSX_PATH.mkdir(parents=True, exist_ok=True)

    jsonl_path = JSONL_PATH/"column_inspection.jsonl"
    output_xlsx = XLSX_PATH/"column_inspection.xlsx"

    path = Path(jsonl_path)
    if not path.exists():
        raise FileNotFoundError(path)

    lines = path.read_text(encoding="utf-8").splitlines()
    if not lines:
        print("JSONL empty")
        return

    # start = max(0, len(lines) - n)
    if last_n is None:
        start = 0
    else:
        start = max(0, len(lines) - last_n)

    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet

    for record_idx, line in enumerate(lines[start:], start=start + 1):
        obj = json.loads(line)

        dataset_name = obj.get("dataset_file_name")
        dataset_path = obj.get("dataset_file_path")

        prefix = f"rec_{record_idx}_"

        # ---------------- metadata ----------------
        ws = wb.create_sheet(prefix + "metadata")
        ws.append(["dataset_file_name", dataset_name])
        ws.append(["dataset_file_path", dataset_path])
        _auto_adjust_column_width(ws)

        # ---------------- column_profiles ----------------
        column_profiles = obj.get("column_profiles", [])
        if column_profiles:
            headers = list(column_profiles[0].keys())
            ws = wb.create_sheet(prefix + "column_profiles")
            _write_table(ws, column_profiles, headers)

        # ---------------- correlation_pairs ----------------
        corr = obj.get("correlation_pairs", [])
        if corr:
            headers = list(corr[0].keys())
            ws = wb.create_sheet(prefix + "correlation_pairs")
            _write_table(ws, corr, headers)

        # ---------------- redundant_features ----------------
        redundant = obj.get("redundant_features", [])
        if redundant:
            headers = list(redundant[0].keys())
            ws = wb.create_sheet(prefix + "redundant_features")
            _write_table(ws, redundant, headers)

        # ---------------- derived_relationships ----------------
        derived = obj.get("derived_relationships", [])
        if derived:
            headers = list(derived[0].keys())
            ws = wb.create_sheet(prefix + "derived_relationships")
            _write_table(ws, derived, headers)

        # ---------------- dependency_graph ----------------
        dep_graph = obj.get("dependency_graph", {})
        if dep_graph:
            ws = wb.create_sheet(prefix + "dependency_graph")
            _write_dependency_graph(ws, dep_graph)

        # ---------------- drop_recommendations ----------------
        drops = obj.get("drop_recommendations", [])
        if drops:
            ws = wb.create_sheet(prefix + "drop_recommendations")
            _write_simple_list(ws, drops, "feature_to_drop")

    output = Path(output_xlsx)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    # print(f"Excel bundle exported → {output}")


# --------------------------------------------------
# run
# --------------------------------------------------
if __name__ == "__main__":
    column_inspection()