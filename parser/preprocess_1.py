# OK
import json
from pathlib import Path
from openpyxl import Workbook
from parser.excel_writer import _auto_adjust_column_width_2


JSONL_FILENAME = "data/preprocesses_1.jsonl"
OUTPUT_FILENAME = "parser/data/xlsx/preprocesses_1.xlsx"


# -----------------------------------------------------
# helpers
# -----------------------------------------------------
def _project_root():
    return Path(__file__).resolve().parents[1]


def _fmt(v):
    if v is None:
        return ""
    if isinstance(v, (list, dict)):
        return json.dumps(v)
    return v


def _write_dict_sheet(wb, sheet_name, data: dict):
    ws = wb.create_sheet(sheet_name[:31])
    ws.append(["key", "value"])
    for k, v in data.items():
        ws.append([k, _fmt(v)])


# def _write_list_of_dicts(wb, sheet_name, rows):
#     ws = wb.create_sheet(sheet_name[:31])
#     if not rows:
#         return

#     headers = sorted({k for r in rows for k in r.keys()})
#     ws.append(headers)

#     for r in rows:
#         ws.append([_fmt(r.get(h)) for h in headers])

def _write_list_of_dicts(wb, sheet_name, rows):
    ws = wb.create_sheet(sheet_name[:31])
    if not rows:
        return

    headers = sorted({k for r in rows for k in r.keys()})
    ws.append(headers)

    for r in rows:

        # detect list fields that need row expansion
        list_fields = {k: v for k, v in r.items() if isinstance(v, list)}

        if not list_fields:
            ws.append([_fmt(r.get(h)) for h in headers])
            continue

        # expand rows for each list element
        max_len = max(len(v) for v in list_fields.values())

        for i in range(max_len):
            new_row = []
            for h in headers:
                val = r.get(h)

                if isinstance(val, list):
                    val = val[i] if i < len(val) else ""
                else:
                    val = val

                new_row.append(_fmt(val))

            ws.append(new_row)


def _write_simple_list(wb, sheet_name, values):
    ws = wb.create_sheet(sheet_name[:31])
    ws.append(["value"])
    for v in values:
        ws.append([_fmt(v)])


# -----------------------------------------------------
# robust JSON loader
# -----------------------------------------------------
def _load_jsonl(path: Path):
    text = path.read_text(encoding="utf-8")

    objs = []
    buffer = ""
    depth = 0
    in_string = False
    escape = False

    for ch in text:
        buffer += ch

        if ch == '"' and not escape:
            in_string = not in_string

        if ch == "\\" and not escape:
            escape = True
            continue
        escape = False

        if in_string:
            continue

        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                objs.append(json.loads(buffer))
                buffer = ""

    return objs


# -----------------------------------------------------
# exporter
# -----------------------------------------------------
def preprocess_1(n: int | None = None):

    root = _project_root()
    data_dir = root

    jsonl_file = data_dir / JSONL_FILENAME
    output_file = data_dir / OUTPUT_FILENAME

    if not jsonl_file.exists():
        raise FileNotFoundError(jsonl_file)

    records = _load_jsonl(jsonl_file)

    if not records:
        print("No records found")
        return

    # -------- last N selection --------
    if n is not None:
        records = records[-n:]

    wb = Workbook()
    wb.remove(wb.active)

    for i, rec in enumerate(records, start=1):
        prefix = f"record_{i}_"

        flat = {}
        nested = {}

        for k, v in rec.items():
            if isinstance(v, (dict, list)):
                nested[k] = v
            else:
                flat[k] = v

        _write_dict_sheet(wb, prefix + "metadata", flat)

        for key, value in nested.items():
            sheet = prefix + key

            if isinstance(value, list):
                if value and isinstance(value[0], dict):
                    _write_list_of_dicts(wb, sheet, value)
                else:
                    _write_simple_list(wb, sheet, value)
            else:
                _write_dict_sheet(wb, sheet, value)

    output_file.parent.mkdir(parents=True, exist_ok=True)
    _auto_adjust_column_width_2(wb)
    wb.save(output_file)

    # print(f"\nExcel written → {output_file}")
    # print(f"Records exported → {len(records)}")


# -----------------------------------------------------
# entry
# -----------------------------------------------------
if __name__ == "__main__":
    preprocess_1()