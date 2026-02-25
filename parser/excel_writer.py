# OK
import json
from openpyxl import Workbook


# def fmt(v):
#     if v is None:
#         return ""
#     if isinstance(v, (dict, list)):
#         return json.dumps(v, indent=2)
#     return v

def fmt(v):
    if v is None:
        return ""

    # list → comma separated string
    if isinstance(v, list):
        return ", ".join(str(x) for x in v)

    # dict → keep JSON (optional, can change later)
    if isinstance(v, dict):
        return json.dumps(v)

    return v


def sheet_table(wb, name, headers, rows):
    ws = wb.create_sheet(name[:31])
    ws.append(headers)
    for r in rows:
        ws.append([fmt(x) for x in r])


def sheet_key_value(wb, name, data: dict):
    ws = wb.create_sheet(name[:31])
    ws.append(["key", "value"])
    for k, v in data.items():
        ws.append([k, fmt(v)])


def sheet_list(wb, name, values, header="value"):
    ws = wb.create_sheet(name[:31])
    ws.append([header])
    for v in values:
        ws.append([fmt(v)])

def _auto_adjust_column_width(ws, min_width=8, max_width=6000, padding=2):
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            try:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))
            except Exception:
                pass

        adjusted_width = max(min_width, min(max_width, max_length + padding))
        ws.column_dimensions[column_letter].width = adjusted_width

def _auto_adjust_column_width_2(wb, min_width=8, max_width=6000, padding=2):
    for ws in wb.worksheets:
        for column_cells in ws.columns:
            if not column_cells:
                continue

            max_length = 0
            column_letter = column_cells[0].column_letter

            for cell in column_cells:
                try:
                    value = "" if cell.value is None else str(cell.value)
                    max_length = max(max_length, len(value))
                except Exception:
                    pass

            adjusted_width = max(min_width, min(max_width, max_length + padding))
            ws.column_dimensions[column_letter].width = adjusted_width